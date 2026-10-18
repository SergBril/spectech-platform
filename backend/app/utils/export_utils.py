"""Утилиты экспорта заявок в PDF и Excel."""

import io
from datetime import datetime
from typing import Any

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT_REGULAR = "app/services/fonts/DejaVuSans.ttf"
FONT_BOLD = "app/services/fonts/DejaVuSans-Bold.ttf"

STATUS_LABELS = {
    "new": "Новая",
    "processing": "В обработке",
    "assigned": "Назначена",
    "confirmed": "Подтверждена",
    "released": "Выпущена",
    "in_progress": "В работе",
    "completed": "Завершена",
    "signed": "Подписана",
    "closed": "Закрыта",
    "cancelled": "Отменена",
    "disputed": "Спорная",
}

PRIORITY_LABELS = {
    "low": "Низкий",
    "normal": "Обычный",
    "high": "Высокий",
    "urgent": "Срочный",
}

COLUMNS = [
    ("№", "id", 10),
    ("Тип работы", "work_type", 30),
    ("Адрес", "object_address", 35),
    ("Период", "period", 52),
    ("Категория", "vehicle_category_name", 28),
    ("Техника", "vehicle_name", 28),
    ("Клиент", "client_name", 28),
    ("Статус", "status_label", 18),
    ("Приоритет", "priority_label", 16),
    ("Сумма", "amount", 18),
]


def _fmt_dt(dt_val):
    """Format datetime to 'DD.MM.YYYY HH:MM' string."""
    if not dt_val:
        return ""
    if isinstance(dt_val, str):
        try:
            d = datetime.fromisoformat(dt_val.replace("Z", "+00:00"))
            return d.strftime("%d.%m.%Y %H:%M")
        except Exception:
            return dt_val[:16].replace("T", " ")
    if hasattr(dt_val, "strftime"):
        return dt_val.strftime("%d.%m.%Y %H:%M")
    return str(dt_val)


def _prepare_rows(requests: list[Any]) -> list[dict]:
    rows = []
    for r in requests:
        data = r if isinstance(r, dict) else {}
        if not isinstance(r, dict):
            for attr in [
                "id",
                "work_type",
                "object_address",
                "planned_start",
                "planned_end",
                "status",
                "priority",
                "estimated_amount",
                "final_amount",
                "vehicle_category_name",
                "assigned_vehicle_name",
                "client_name",
            ]:
                data[attr] = getattr(r, attr, None)

        ps = data.get("planned_start")
        pe = data.get("planned_end")
        start_str = _fmt_dt(ps)
        end_str = _fmt_dt(pe)
        if start_str and end_str:
            period = start_str + " - " + end_str
        elif start_str:
            period = start_str
        else:
            period = "-"

        status = data.get("status", "")
        priority = data.get("priority", "")
        amount = data.get("estimated_amount") or data.get("final_amount")

        rows.append(
            {
                "id": str(data.get("id", "")),
                "work_type": str(data.get("work_type", "") or "-"),
                "object_address": str(data.get("object_address", "") or "-"),
                "period": period,
                "vehicle_category_name": str(
                    data.get("vehicle_category_name", "") or "-"
                ),
                "vehicle_name": str(data.get("assigned_vehicle_name", "") or "-"),
                "client_name": str(data.get("client_name", "") or "-"),
                "status_label": STATUS_LABELS.get(status, str(status)),
                "priority_label": PRIORITY_LABELS.get(priority, str(priority)),
                "amount": f"{float(amount):,.0f} RUB" if amount else "-",
            }
        )
    return rows


def export_requests_pdf(requests: list[Any], title: str = "Заявки") -> bytes:
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.add_font("DejaVu", "", FONT_REGULAR)
    pdf.add_font("DejaVu", "B", FONT_BOLD)
    pdf.set_auto_page_break(auto=True, margin=10)
    pdf.add_page()

    # Title
    pdf.set_font("DejaVu", "B", 14)
    pdf.cell(0, 10, title, ln=True, align="C")
    pdf.set_font("DejaVu", "", 8)
    pdf.cell(
        0, 6, "Data: " + datetime.now().strftime("%d.%m.%Y %H:%M"), ln=True, align="C"
    )
    pdf.ln(3)

    # Header
    pdf.set_font("DejaVu", "B", 6)
    pdf.set_fill_color(230, 230, 230)
    for label, _, width in COLUMNS:
        pdf.cell(width, 7, label, border=1, fill=True, align="C")
    pdf.ln()

    # Data rows
    pdf.set_font("DejaVu", "", 5)
    rows = _prepare_rows(requests)
    for row in rows:
        for _, key, width in COLUMNS:
            val = str(row.get(key, ""))
            # Truncate to prevent overflow
            max_chars = int(width / 1.8)
            if len(val) > max_chars:
                val = val[:max_chars]
            align = "R" if key == "amount" else "L"
            pdf.cell(width, 5, val, border=1, align=align)
        pdf.ln()

    return pdf.output()


def export_requests_excel(requests: list[Any], title: str = "Заявки") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]

    header_font = Font(bold=True, size=10)
    header_fill = PatternFill(
        start_color="E0E0E0", end_color="E0E0E0", fill_type="solid"
    )
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Title (only first cell, no merge to avoid duplication)
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(bold=True, size=14)

    ws.cell(
        row=2, column=1, value="Data: " + datetime.now().strftime("%d.%m.%Y %H:%M")
    ).font = Font(size=9)

    # Headers at row 4
    for col_idx, (label, _, _) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=4, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    rows = _prepare_rows(requests)
    for row_idx, row in enumerate(rows, 5):
        for col_idx, (_, key, _) in enumerate(COLUMNS, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row.get(key, ""))
            cell.border = thin_border

    # Column widths
    from openpyxl.utils import get_column_letter

    for col_idx, (_, _, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width / 2.0

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
