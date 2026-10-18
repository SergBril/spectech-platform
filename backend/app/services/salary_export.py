"""Excel экспорт зарплат."""

import logging
from io import BytesIO

logger = logging.getLogger("spectech.salary_export")

SALARY_TYPE_LABELS = {
    "fixed_shift": "Фикс/смена",
    "hourly": "Почасовая",
    "per_km": "За пробег",
    "shift_percent": "% от смены",
    "combo": "Комбинация",
}


def export_salaries_to_excel(
    records: list, employee_names: dict[int, str] | None = None
) -> bytes:
    """
    Экспортирует записи зарплат в Excel.

    Args:
        records: список SalaryRecord
        employee_names: dict {contractor_id: name}

    Returns:
        bytes: xlsx содержимое
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError:
        raise ImportError(
            "openpyxl is required for Excel export. Install: pip install openpyxl"
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "Зарплаты"

    # Стили
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(
        start_color="1E40AF", end_color="1E40AF", fill_type="solid"
    )
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    money_format = "#,##0.00"

    # Заголовки
    headers = [
        "№",
        "Сотрудник ID",
        "Имя",
        "Период с",
        "Период по",
        "Тип оплаты",
        "Смены",
        "Часы",
        "Пробег (км)",
        "Начислено",
        "Удержания",
        "К выплате",
        "Статус",
        "Дата расчёта",
        "Формула",
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Данные
    for i, r in enumerate(records, 1):
        name = ""
        if employee_names and r.employee_contractor_id in employee_names:
            name = employee_names[r.employee_contractor_id]

        formula = ""
        if r.calculation_details and isinstance(r.calculation_details, dict):
            formula = r.calculation_details.get("formula", "")

        status_label = {
            "draft": "Черновик",
            "approved": "Утверждено",
            "paid": "Выплачено",
        }.get(r.status, r.status)

        row = [
            i,
            r.employee_contractor_id,
            name,
            r.period_start.isoformat() if r.period_start else "",
            r.period_end.isoformat() if r.period_end else "",
            SALARY_TYPE_LABELS.get(r.salary_type, r.salary_type),
            r.shifts_count,
            float(r.total_hours or 0),
            float(r.total_km or 0),
            float(r.gross_amount or 0),
            float(r.deductions or 0),
            float(r.net_amount or 0),
            status_label,
            r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else "",
            formula,
        ]
        ws.append(row)

        # Форматирование денежных колонок
        last_row = ws.max_row
        for col_idx in [10, 11, 12]:  # Начислено, Удержания, К выплате
            cell = ws.cell(row=last_row, column=col_idx)
            cell.number_format = money_format
            cell.border = thin_border

        # Границы для всех ячеек
        for cell in ws[last_row]:
            cell.border = thin_border

    # Ширина колонок
    column_widths = [5, 10, 25, 12, 12, 15, 8, 8, 10, 14, 12, 14, 12, 18, 30]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = width

    # Автофильтр
    ws.auto_filter.ref = ws.dimensions

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()
